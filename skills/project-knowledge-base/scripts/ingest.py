#!/usr/bin/env python3
"""ingest — turn a folder of project documents into text the knowledge-base workflow can read and cite.

    uv run --with pypdf --with python-docx --with openpyxl --with python-pptx \
        python ingest.py <sources-dir> <kb-dir>

Reads <sources-dir> (never writes there) and writes, under <kb-dir>/_sources/:
  manifest.json   one entry per file: id (S01…), path, type, sha256, words, title, outline (heading → line), status,
                  and change status against the previous run (new / changed / unchanged / removed)
  text/<id>-<slug>.md   extracted text, with <!-- page N --> / <!-- sheet X --> / <!-- slide N --> markers
  inventory.md    the same as a table for people

Formats: .md .markdown .txt .rst .csv .json .yaml .yml .html .htm (stdlib); .pdf (pypdf); .docx (python-docx);
.xlsx (openpyxl); .pptx (python-pptx). Images are listed with status "image" for the model to view directly.
A PDF with no extractable text is "needs-ocr". Missing optional libraries make that format "unsupported", not a crash.
IDs are stable across runs: a file keeps its id while its path stays the same.
"""

from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import re
import sys
from datetime import UTC, datetime
from pathlib import Path

TEXT = {".md", ".markdown", ".txt", ".rst", ".json", ".yaml", ".yml"}
IMAGES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg", ".bmp"}
SKIP_DIRS = {".git", "node_modules", "__pycache__", ".venv", "venv"}
HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")


def slug(text: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return s[:50] or "doc"


def read_text_file(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return path.read_text(encoding=encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def from_html(path: Path) -> str:
    raw = read_text_file(path)
    raw = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", "", raw)
    raw = re.sub(r"(?i)<h([1-6])[^>]*>", lambda m: "\n" + "#" * int(m.group(1)) + " ", raw)
    raw = re.sub(r"(?i)<(br|/p|/div|/li|/tr|/h[1-6])[^>]*>", "\n", raw)
    raw = re.sub(r"<[^>]+>", "", raw)
    return html.unescape(re.sub(r"\n{3,}", "\n\n", raw))


def from_csv(path: Path) -> str:
    rows = list(csv.reader(io.StringIO(read_text_file(path))))
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    lines = ["| " + " | ".join(c.replace("|", "\\|") for c in r + [""] * (width - len(r))) + " |" for r in rows]
    lines.insert(1, "|" + "---|" * width)
    return "\n".join(lines)


def from_pdf(path: Path) -> tuple[str, str]:
    try:
        from pypdf import PdfReader
    except ImportError:
        return "", "unsupported"
    reader = PdfReader(str(path))
    parts = []
    for number, page in enumerate(reader.pages, start=1):
        parts.append(f"<!-- page {number} -->\n{(page.extract_text() or '').strip()}")
    text = "\n\n".join(parts)
    has_text = len(re.sub(r"<!-- page \d+ -->|\s", "", text)) > 20 * max(1, len(reader.pages)) // 4
    return text, "ok" if has_text else "needs-ocr"


def from_docx(path: Path) -> tuple[str, str]:
    try:
        import docx
    except ImportError:
        return "", "unsupported"
    document = docx.Document(str(path))
    lines = []
    for block in document.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = docx.text.paragraph.Paragraph(block, document)
            style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
            text = paragraph.text.strip()
            if not text:
                continue
            level = re.match(r"heading (\d)", style)
            if level:
                lines.append("#" * int(level.group(1)) + " " + text)
            elif style == "title":
                lines.append("# " + text)
            elif "list" in style:
                lines.append("- " + text)
            else:
                lines.append(text)
        elif tag == "tbl":
            table = docx.table.Table(block, document)
            rows = [[cell.text.strip().replace("|", "\\|").replace("\n", " ") for cell in row.cells] for row in table.rows]
            if rows:
                lines.append("| " + " | ".join(rows[0]) + " |")
                lines.append("|" + "---|" * len(rows[0]))
                lines.extend("| " + " | ".join(r) + " |" for r in rows[1:])
        lines.append("")
    return "\n".join(lines), "ok"


def from_xlsx(path: Path) -> tuple[str, str]:
    try:
        import openpyxl
    except ImportError:
        return "", "unsupported"
    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for sheet in workbook.worksheets:
        rows = [["" if v is None else str(v).replace("|", "\\|").replace("\n", " ") for v in row]
                for row in sheet.iter_rows(values_only=True)]
        rows = [r for r in rows if any(c.strip() for c in r)]
        parts.append(f"<!-- sheet {sheet.title} -->\n## {sheet.title}\n")
        if rows:
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            parts.append("| " + " | ".join(rows[0]) + " |\n|" + "---|" * width)
            parts.extend("| " + " | ".join(r) + " |" for r in rows[1:])
        parts.append("")
    return "\n".join(parts), "ok"


def from_pptx(path: Path) -> tuple[str, str]:
    try:
        import pptx
    except ImportError:
        return "", "unsupported"
    deck = pptx.Presentation(str(path))
    parts = []
    for number, slide in enumerate(deck.slides, start=1):
        parts.append(f"<!-- slide {number} -->")
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        parts.append(f"## Slide {number}{': ' + title if title else ''}")
        for shape in slide.shapes:
            if shape.has_text_frame and shape is not slide.shapes.title:
                parts.extend(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"> Notes: {notes}")
        parts.append("")
    return "\n".join(parts), "ok"


def extract(path: Path) -> tuple[str, str, str]:
    """Return (kind, text, status)."""
    ext = path.suffix.lower()
    if ext in IMAGES:
        return "image", "", "image"
    if ext in TEXT:
        return ext.lstrip("."), read_text_file(path), "ok"
    if ext in {".html", ".htm"}:
        return "html", from_html(path), "ok"
    if ext == ".csv":
        return "csv", from_csv(path), "ok"
    handlers = {".pdf": from_pdf, ".docx": from_docx, ".xlsx": from_xlsx, ".pptx": from_pptx}
    if ext in handlers:
        try:
            text, status = handlers[ext](path)
        except Exception as error:  # a corrupt file must not stop the whole run
            return ext.lstrip("."), "", f"error: {type(error).__name__}: {error}"[:200]
        return ext.lstrip("."), text, status
    return ext.lstrip(".") or "unknown", "", "unsupported"


def outline(text: str) -> list[dict[str, object]]:
    found = []
    in_fence = False
    for number, line in enumerate(text.splitlines(), start=1):
        if line.lstrip().startswith("```"):
            in_fence = not in_fence
        if in_fence:
            continue
        match = HEADING.match(line)
        if match:
            found.append({"level": len(match.group(1)), "title": match.group(2).strip(), "line": number})
    return found


def main() -> int:
    if len(sys.argv) != 3:
        print(__doc__)
        return 2
    sources, kb = Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve()
    if not sources.is_dir():
        print(f"not a directory: {sources}")
        return 2
    if kb == sources or sources in kb.parents:
        print("the knowledge base must not be written inside the sources folder")
        return 2
    out = kb / "_sources"
    text_dir = out / "text"
    text_dir.mkdir(parents=True, exist_ok=True)

    manifest_path = out / "manifest.json"
    previous = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else {"files": []}
    previous_by_path = {f["path"]: f for f in previous.get("files", [])}
    used_ids = {f["id"] for f in previous.get("files", [])}
    next_number = max([int(i[1:]) for i in used_ids] + [0]) + 1

    files = sorted(
        p for p in sources.rglob("*")
        if p.is_file() and not any(part in SKIP_DIRS or part.startswith(".") for part in p.relative_to(sources).parts)
    )
    entries = []
    for path in files:
        rel = path.relative_to(sources).as_posix()
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        old = previous_by_path.get(rel)
        if old:
            doc_id = old["id"]
        else:
            doc_id = f"S{next_number:02d}"
            next_number += 1
        kind, text, status = extract(path)
        heads = outline(text)
        title = next((h["title"] for h in heads if h["level"] == 1), None) or path.stem.replace("_", " ")
        text_rel = None
        if text.strip():
            text_file = text_dir / f"{doc_id}-{slug(path.stem)}.md"
            header = f"<!-- {doc_id}: {rel} — extracted {datetime.now(UTC):%Y-%m-%d}; cite as [{doc_id} L<line>] -->\n"
            text_file.write_text(header + text, encoding="utf-8")
            text_rel = text_file.relative_to(kb).as_posix()
        change = "new" if not old else ("unchanged" if old.get("sha256") == digest else "changed")
        entries.append({
            "id": doc_id, "path": rel, "type": kind, "bytes": path.stat().st_size, "sha256": digest,
            "modified": datetime.fromtimestamp(path.stat().st_mtime, UTC).strftime("%Y-%m-%d"),
            "words": len(text.split()), "title": title, "status": status, "change": change,
            "text": text_rel, "outline": heads,
        })
    removed = [dict(f, change="removed") for p, f in previous_by_path.items() if p not in {e["path"] for e in entries}]
    manifest = {"sources": str(sources), "generated": datetime.now(UTC).strftime("%Y-%m-%dT%H:%MZ"),
                "files": entries + removed}
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    rows = ["| ID | Source | Type | Words | Modified | Status | Change |", "|---|---|---|---:|---|---|---|"]
    for e in manifest["files"]:
        link = f"[{e['path']}]({e['text'].removeprefix('_sources/')})" if e.get("text") else e["path"]
        rows.append(f"| {e['id']} | {link} | {e['type']} | {e.get('words', 0)} | {e.get('modified', '')} | {e['status']} | {e['change']} |")
    (out / "inventory.md").write_text(
        f"# Source inventory\n\nFrom `{sources}` on {manifest['generated']}. Generated by ingest.py — do not edit.\n\n"
        + "\n".join(rows) + "\n", encoding="utf-8")

    counts: dict[str, int] = {}
    for e in manifest["files"]:
        key = e["status"] if e["change"] != "removed" else "removed"
        counts[key] = counts.get(key, 0) + 1
    changes = {c: sum(1 for e in manifest["files"] if e["change"] == c) for c in ("new", "changed", "unchanged", "removed")}
    print(f"{len(entries)} files | status {counts} | change {changes}")
    print(f"words extracted: {sum(e['words'] for e in entries)} | manifest: {manifest_path}")
    for e in manifest["files"]:
        if e["status"] not in ("ok", "image"):
            print(f"  {e['id']} {e['path']}: {e['status']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
